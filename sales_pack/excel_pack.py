"""Write the weekly sales pack workbook.

Every metric in the pack is an Excel formula over the Data and Quotas sheets, driven by two
inputs on the Summary sheet (as-of date and quota growth). Change either input in Excel and the
whole pack recalculates, so it can be reviewed for any week without rerunning Python.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from . import __version__
from .periods import Period

FONT = "Arial"
NAVY = "1F3A5F"
BAND = "EEF2F7"
GREY = "595959"
INPUT_FILL = "FFFF00"  # yellow: cells the reader is meant to change
INPUT_FONT = "0000FF"  # blue: hard-coded inputs
GOOD, WARN, BAD = "C6EFCE", "FFEB9C", "FFC7CE"

MONEY = '$#,##0;($#,##0);"-"'
MONEY2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
PTS = '0.0" pts";(0.0" pts");"-"'
INT = '#,##0;(#,##0);"-"'
DATE = "mmm d, yyyy"

THIN = Side(style="thin", color="8EA0B8")

# (header, field, number format) for the Data sheet. Formulas reference these columns
# through the d_* defined names, so keep the two lists in step.
DATA_COLUMNS = [
    ("Order ID", "order_id", None),
    ("Order Date", "order_date", DATE),
    ("Week Ending", "week_ending", DATE),
    ("Month", "month", "mmm yyyy"),
    ("Region", "region", None),
    ("State", "state", None),
    ("Segment", "segment", None),
    ("Category", "category", None),
    ("Sub-Category", "sub_category", None),
    ("Sales", "sales", MONEY2),
    ("Quantity", "quantity", INT),
    ("Profit", "profit", MONEY2),
    ("Order Count", "order_count", INT),
]
DATA_NAMES = {
    "d_date": "order_date", "d_region": "region", "d_category": "category",
    "d_subcat": "sub_category", "d_sales": "sales", "d_orders": "order_count",
}
REGIONAL_HEADERS = ["Region", "Revenue", "Quota to date", "Attainment", "Gap to quota",
                    "Orders", "Avg order value", "PY revenue", "Growth vs PY"]


@dataclass
class RunInfo:
    source: str
    rows_loaded: int
    rows_dropped: int
    first_date: pd.Timestamp
    last_date: pd.Timestamp
    as_of: pd.Timestamp
    growth: float
    quota_method: str
    trend_weeks: int


def write_pack(path: Path, sales: pd.DataFrame, quotas: pd.DataFrame, periods: list[Period],
               regions: list[str], categories: list[str], subcategories: pd.DataFrame,
               callouts: list[str], info: RunInfo) -> Path:
    wb = Workbook()
    _use_arial_by_default(wb)
    summary = wb.active
    summary.title = "Summary"
    regional = wb.create_sheet("Regional")
    trend = wb.create_sheet("Weekly Trend")
    mix = wb.create_sheet("Sales Mix")
    quota_ws = wb.create_sheet("Quotas")
    data = wb.create_sheet("Data")
    notes = wb.create_sheet("Definitions")

    _write_data(wb, data, sales)
    _write_controls(wb, summary, info)
    _write_quotas(wb, quota_ws, quotas)
    layout = _write_regional(regional, periods, regions)
    _write_trend(trend, regions, info.trend_weeks)
    _write_mix(mix, categories, regions, subcategories)
    _write_summary(summary, layout, regions, callouts, info)
    _write_definitions(notes, info)

    for ws in (summary, regional, trend, mix, quota_ws, notes):
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "portrait" if ws is summary else "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True  # Excel computes every formula when the file opens
    wb.save(path)
    return path


# --------------------------------------------------------------------------- helpers

def _use_arial_by_default(wb: Workbook) -> None:
    """Make Arial the workbook default so unstyled cells (e.g. the Data sheet) match the report."""
    try:
        default = Font(name=FONT, size=10)
        wb._fonts[0] = default  # openpyxl keeps the default font first in its style table
        wb._named_styles["Normal"].font = default
    except (AttributeError, IndexError, KeyError, TypeError):
        pass  # cosmetic only: fall back to openpyxl's default font


def _font(bold=False, color="000000", size=10, italic=False) -> Font:
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def _define(wb: Workbook, name: str, ref: str) -> None:
    dn = DefinedName(name, attr_text=ref)
    try:
        wb.defined_names[name] = dn  # openpyxl 3.1+
    except TypeError:
        wb.defined_names.append(dn)  # openpyxl 3.0


def _title(ws, text: str, subtitle: str | None = None) -> None:
    ws["A1"] = text
    ws["A1"].font = _font(bold=True, color=NAVY, size=14)
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = _font(italic=True, color=GREY, size=9)


def _header(ws, row: int, labels: list[str], col: int = 1) -> None:
    for i, label in enumerate(labels):
        cell = ws.cell(row=row, column=col + i, value=label)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _section(ws, row: int, text: str, width: int, col: int = 1) -> None:
    for c in range(col, col + width):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = _font(bold=True, color="FFFFFF")
    ws.cell(row=row, column=col, value=text)


def _put(ws, row: int, col: int, value, fmt: str | None = None, bold: bool = False):
    cell = ws.cell(row=row, column=col, value=value)
    if fmt:
        cell.number_format = fmt
    if bold:
        cell.font = _font(bold=True)
    return cell


def _total_row(ws, row: int, first_col: int, last_col: int) -> None:
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = _font(bold=True)
        cell.fill = PatternFill("solid", fgColor=BAND)
        cell.border = Border(top=THIN)


def _attainment_colors(ws, cell_range: str) -> None:
    """Green at or above 100% of quota to date, amber from 90%, red below 90%."""
    ws.conditional_formatting.add(cell_range, CellIsRule(
        operator="greaterThanOrEqual", formula=["1"], fill=PatternFill("solid", fgColor=GOOD)))
    ws.conditional_formatting.add(cell_range, CellIsRule(
        operator="between", formula=["0.9", "0.99999"], fill=PatternFill("solid", fgColor=WARN)))
    ws.conditional_formatting.add(cell_range, CellIsRule(
        operator="lessThan", formula=["0.9"], fill=PatternFill("solid", fgColor=BAD)))


def _widths(ws, widths: dict[str, float]) -> None:
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def _excel_value(v):
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if hasattr(v, "item"):  # numpy scalar -> Python scalar
        return v.item()
    return v


def _in_window(start: str, end: str = "AsOf") -> str:
    """SUMIFS criteria pair for d_date between two dates (names or expressions)."""
    return f'd_date,">="&{start},d_date,"<="&{end}'


# --------------------------------------------------------------------------- sheets

def _write_data(wb: Workbook, ws, sales: pd.DataFrame) -> None:
    fields = [field for _, field, _ in DATA_COLUMNS]
    _header(ws, 1, [h for h, _, _ in DATA_COLUMNS])
    for values in sales[fields].itertuples(index=False, name=None):
        ws.append([_excel_value(v) for v in values])
    last = max(len(sales) + 1, 2)
    for idx, (_, field, fmt) in enumerate(DATA_COLUMNS, start=1):
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, max_row=last, min_col=idx, max_col=idx):
                cell.number_format = fmt
        ws.column_dimensions[get_column_letter(idx)].width = 18 if field in ("order_id", "sub_category") else 13
    for name, field in DATA_NAMES.items():
        col = get_column_letter(fields.index(field) + 1)
        _define(wb, name, f"Data!${col}$2:${col}${last}")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{last}"


def _write_controls(wb: Workbook, ws, info: RunInfo) -> None:
    _title(ws, "Regional Sales Pack",
           f"Weekly review · built {datetime.now():%b %d, %Y %H:%M} by sales_pack v{__version__}")
    _section(ws, 4, "Controls", 2)
    rows = [
        ("As-of date (week ending)", info.as_of.to_pydatetime(), DATE, "AsOf", True),
        ("Quota growth vs last year", info.growth, PCT, "Growth", True),
        ("Week start", "=AsOf-6", DATE, "WeekStart", False),
        ("Month start", "=DATE(YEAR(AsOf),MONTH(AsOf),1)", DATE, "MonthStart", False),
        ("Quarter start", "=DATE(YEAR(AsOf),INT((MONTH(AsOf)-1)/3)*3+1,1)", DATE, "QtrStart", False),
        ("Year start", "=DATE(YEAR(AsOf),1,1)", DATE, "YearStart", False),
    ]
    for r, (label, value, fmt, name, is_input) in enumerate(rows, start=5):
        ws.cell(row=r, column=1, value=label)
        cell = _put(ws, r, 2, value, fmt)
        cell.alignment = Alignment(horizontal="right")
        if is_input:
            cell.font = _font(bold=True, color=INPUT_FONT)
            cell.fill = PatternFill("solid", fgColor=INPUT_FILL)
        _define(wb, name, f"Summary!$B${r}")
    ws.cell(row=11, column=1, value="Change the yellow cells to rerun the whole pack for another week or "
                                    "growth target.").font = _font(italic=True, color=GREY, size=9)


def _write_quotas(wb: Workbook, ws, quotas: pd.DataFrame) -> None:
    _title(ws, "Monthly quotas by region",
           "Quota = same month last year × (1 + growth on the Summary sheet). Blue quotas come from the "
           "override file. The four right-hand columns pro-rate each month by the days inside each window.")
    headers = ["Region", "Month", "Month end", "Days", "PY revenue", "Quota",
               "Week quota", "MTD quota", "QTD quota", "YTD quota", "Source"]
    _header(ws, 4, headers)
    first = 5
    for r, q in enumerate(quotas.itertuples(index=False), start=first):
        ws.cell(row=r, column=1, value=q.region)
        _put(ws, r, 2, q.month.to_pydatetime(), "mmm yyyy")
        _put(ws, r, 3, f"=EOMONTH(B{r},0)", DATE)
        _put(ws, r, 4, f"=DAY(C{r})")
        _put(ws, r, 5, f'=SUMIFS(d_sales,d_region,A{r},{_in_window(f"EDATE(B{r},-12)", f"EOMONTH(EDATE(B{r},-12),0)")})',
             MONEY)
        if q.source == "override":
            _put(ws, r, 6, float(q.quota), MONEY).font = _font(color=INPUT_FONT)
        else:
            _put(ws, r, 6, f"=E{r}*(1+Growth)", MONEY)
        for col, start in zip(range(7, 11), ("WeekStart", "MonthStart", "QtrStart", "YearStart")):
            _put(ws, r, col, f"=IF(D{r}=0,0,F{r}*MAX(0,MIN(AsOf,C{r})-MAX({start},B{r})+1)/D{r})", MONEY)
        ws.cell(row=r, column=11, value=q.source)
    last = max(first, first + len(quotas) - 1)
    for name, col in (("q_region", "A"), ("q_week", "G"), ("q_mtd", "H"), ("q_qtd", "I"), ("q_ytd", "J")):
        _define(wb, name, f"Quotas!${col}${first}:${col}${last}")
    ws.freeze_panes = "A5"
    _widths(ws, {"A": 14, "B": 11, "C": 13, "D": 7, "E": 13, "F": 13,
                 "G": 12, "H": 12, "I": 12, "J": 12, "K": 10})


def _write_regional(ws, periods: list[Period], regions: list[str]) -> dict:
    _title(ws, "Regional performance",
           "Revenue and orders come from the Data sheet. Quota to date is pro-rated by day from the Quotas sheet.")
    layout = {}
    row = 4
    for p in periods:
        start = p.excel_start_name
        ws.cell(row=row, column=1, value=p.label).font = _font(bold=True, color=NAVY, size=11)
        _put(ws, row, 2, f"={start}", DATE).font = _font(color=GREY)
        _put(ws, row, 3, "=AsOf", DATE).font = _font(color=GREY)
        _header(ws, row + 1, REGIONAL_HEADERS)
        first = row + 2
        for j, region in enumerate(regions):
            r = first + j
            ws.cell(row=r, column=1, value=region)
            _put(ws, r, 2, f"=SUMIFS(d_sales,d_region,$A{r},{_in_window(start)})")
            _put(ws, r, 3, f"=SUMIFS(q_{p.key},q_region,$A{r})")
            _put(ws, r, 6, f"=SUMIFS(d_orders,d_region,$A{r},{_in_window(start)})")
            _put(ws, r, 8, f"=SUMIFS(d_sales,d_region,$A{r},{_in_window(*p.excel_py_window)})")
            _regional_ratios(ws, r)
        total = first + len(regions)
        ws.cell(row=total, column=1, value="Total")
        for col in (2, 3, 6, 8):
            letter = get_column_letter(col)
            ws.cell(row=total, column=col, value=f"=SUM({letter}{first}:{letter}{total - 1})")
        _regional_ratios(ws, total)
        for r in range(first, total + 1):
            for col, fmt in zip(range(2, 10), (MONEY, MONEY, PCT, MONEY, INT, MONEY, MONEY, PCT)):
                ws.cell(row=r, column=col).number_format = fmt
        _total_row(ws, total, 1, 9)
        _attainment_colors(ws, f"D{first}:D{total}")
        layout[p.key] = {"total": total, "rows": {reg: first + j for j, reg in enumerate(regions)}}
        row = total + 3
    _widths(ws, {"A": 18, **{get_column_letter(c): 14 for c in range(2, 10)}})
    return layout


def _regional_ratios(ws, r: int) -> None:
    ws.cell(row=r, column=4, value=f'=IF(C{r}=0,"",B{r}/C{r})')
    ws.cell(row=r, column=5, value=f'=IF(C{r}=0,"",B{r}-C{r})')
    ws.cell(row=r, column=7, value=f'=IF(F{r}=0,"",B{r}/F{r})')
    ws.cell(row=r, column=9, value=f'=IF(H{r}=0,"",B{r}/H{r}-1)')


def _write_trend(ws, regions: list[str], weeks: int) -> None:
    _title(ws, f"Weekly trend (last {weeks} weeks)",
           "Weeks run Monday to Sunday. Last year's week is 52 weeks earlier, so weekdays line up.")
    hdr = 4
    _header(ws, hdr, ["Week ending", *regions, "Total", "Same week last year", "Growth vs PY"])
    n = len(regions)
    total_col, py_col, growth_col = n + 2, n + 3, n + 4
    tl, pl, last_region = get_column_letter(total_col), get_column_letter(py_col), get_column_letter(n + 1)
    first = hdr + 1
    for i in range(weeks):
        r = first + i
        back = 7 * (weeks - 1 - i)
        _put(ws, r, 1, f"=AsOf-{back}" if back else "=AsOf", DATE)
        for j in range(n):
            col = get_column_letter(2 + j)
            _put(ws, r, 2 + j, f"=SUMIFS(d_sales,d_region,{col}${hdr},{_in_window(f'($A{r}-6)', f'$A{r}')})", MONEY)
        _put(ws, r, total_col, f"=SUM(B{r}:{last_region}{r})", MONEY, bold=True)
        _put(ws, r, py_col, f"=SUMIFS(d_sales,{_in_window(f'($A{r}-370)', f'($A{r}-364)')})", MONEY)
        _put(ws, r, growth_col, f'=IF({pl}{r}=0,"",{tl}{r}/{pl}{r}-1)', PCT)
    last = first + weeks - 1

    chart = LineChart()
    chart.title = "Weekly revenue by region"
    chart.height, chart.width = 8.5, 24
    chart.y_axis.title = "Revenue"
    chart.y_axis.number_format = "$#,##0"
    chart.x_axis.number_format = "mmm d"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.add_data(Reference(ws, min_col=2, min_row=hdr, max_col=n + 1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    for series in chart.series:
        series.smooth = False
    ws.add_chart(chart, f"A{last + 3}")
    ws.freeze_panes = f"B{first}"
    _widths(ws, {"A": 14, **{get_column_letter(c): 14 for c in range(2, growth_col + 1)}})


def _write_mix(ws, categories: list[str], regions: list[str], subcategories: pd.DataFrame) -> None:
    _title(ws, "Sales mix (quarter to date)",
           "Share of revenue by category and sub-category, compared with the same dates last year.")
    qtd = _in_window("QtrStart")
    pyq = _in_window("EDATE(QtrStart,-12)", "EDATE(AsOf,-12)")
    nc, nr = len(categories), len(regions)

    # Block 1: category mix vs last year
    top = 4
    _header(ws, top, ["Category", "QTD revenue", "Share", "PY QTD revenue", "PY share", "Mix shift", "Growth vs PY"])
    first, total = top + 1, top + 1 + nc
    for i, cat in enumerate(categories):
        r = first + i
        ws.cell(row=r, column=1, value=cat)
        _put(ws, r, 2, f"=SUMIFS(d_sales,d_category,$A{r},{qtd})", MONEY)
        _put(ws, r, 3, f'=IF(B${total}=0,"",B{r}/B${total})', PCT)
        _put(ws, r, 4, f"=SUMIFS(d_sales,d_category,$A{r},{pyq})", MONEY)
        _put(ws, r, 5, f'=IF(D${total}=0,"",D{r}/D${total})', PCT)
        _put(ws, r, 6, f'=IF(OR(C{r}="",E{r}=""),"",(C{r}-E{r})*100)', PTS)
        _put(ws, r, 7, f'=IF(D{r}=0,"",B{r}/D{r}-1)', PCT)
    ws.cell(row=total, column=1, value="Total")
    _put(ws, total, 2, f"=SUM(B{first}:B{total - 1})", MONEY)
    _put(ws, total, 3, f'=IF(B{total}=0,"",1)', PCT)
    _put(ws, total, 4, f"=SUM(D{first}:D{total - 1})", MONEY)
    _put(ws, total, 5, f'=IF(D{total}=0,"",1)', PCT)
    _put(ws, total, 7, f'=IF(D{total}=0,"",B{total}/D{total}-1)', PCT)
    _total_row(ws, total, 1, 7)

    # Block 2: revenue by category and region (feeds the chart)
    top2 = total + 3
    ws.cell(row=top2 - 1, column=1, value="QTD revenue by category and region").font = _font(bold=True, color=NAVY)
    _header(ws, top2, ["Category", *regions, "Total"])
    first2, total2 = top2 + 1, top2 + 1 + nc
    for i, cat in enumerate(categories):
        r = first2 + i
        ws.cell(row=r, column=1, value=cat)
        for j in range(nr):
            col = get_column_letter(2 + j)
            _put(ws, r, 2 + j, f"=SUMIFS(d_sales,d_region,{col}${top2},d_category,$A{r},{qtd})", MONEY)
        _put(ws, r, nr + 2, f"=SUM(B{r}:{get_column_letter(nr + 1)}{r})", MONEY)
    ws.cell(row=total2, column=1, value="Total")
    for c in range(2, nr + 3):
        col = get_column_letter(c)
        _put(ws, total2, c, f"=SUM({col}{first2}:{col}{total2 - 1})", MONEY)
    _total_row(ws, total2, 1, nr + 2)

    # Block 3: each category's share of the region's revenue
    top3 = total2 + 3
    ws.cell(row=top3 - 1, column=1, value="Category share of each region's revenue").font = _font(bold=True, color=NAVY)
    _header(ws, top3, ["Category", *regions, "Total"])
    for i, cat in enumerate(categories):
        r, src = top3 + 1 + i, first2 + i
        ws.cell(row=r, column=1, value=cat)
        for c in range(2, nr + 3):
            col = get_column_letter(c)
            _put(ws, r, c, f'=IF({col}${total2}=0,"",{col}{src}/{col}${total2})', PCT)

    # Block 4: sub-category drivers (sorted by revenue as of the refresh run)
    top4 = top3 + nc + 3
    ws.cell(row=top4 - 1, column=1, value="Sub-category drivers").font = _font(bold=True, color=NAVY)
    _header(ws, top4, ["Sub-category", "Category", "QTD revenue", "PY QTD revenue", "Change", "Growth vs PY"])
    for i, row in enumerate(subcategories.itertuples(index=False), start=top4 + 1):
        ws.cell(row=i, column=1, value=row.sub_category)
        ws.cell(row=i, column=2, value=row.category)
        _put(ws, i, 3, f"=SUMIFS(d_sales,d_subcat,$A{i},d_category,$B{i},{qtd})", MONEY)
        _put(ws, i, 4, f"=SUMIFS(d_sales,d_subcat,$A{i},d_category,$B{i},{pyq})", MONEY)
        _put(ws, i, 5, f"=C{i}-D{i}", MONEY)
        _put(ws, i, 6, f'=IF(D{i}=0,"",C{i}/D{i}-1)', PCT)
    last4 = top4 + len(subcategories)
    if len(subcategories):
        ws.conditional_formatting.add(f"E{top4 + 1}:E{last4}", CellIsRule(
            operator="lessThan", formula=["0"], font=Font(name=FONT, color="C00000")))

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "percentStacked"
    chart.overlap = 100
    chart.title = "Category share of revenue by region (QTD)"
    chart.height, chart.width = 8.5, 16
    chart.y_axis.number_format = "0%"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.add_data(Reference(ws, min_col=1, min_row=first2, max_col=nr + 1, max_row=total2 - 1),
                   from_rows=True, titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=top2, max_col=nr + 1, max_row=top2))
    ws.add_chart(chart, "I4")
    _widths(ws, {"A": 22, **{get_column_letter(c): 14 for c in range(2, max(8, nr + 3) + 1)}})


def _write_summary(ws, layout: dict, regions: list[str], callouts: list[str], info: RunInfo) -> None:
    week_t, qtd_t, ytd_t = layout["week"]["total"], layout["qtd"]["total"], layout["ytd"]["total"]
    _section(ws, 13, "Headline KPIs", 2)
    kpis = [
        ("Revenue, week", f"=Regional!B{week_t}", MONEY),
        ("Revenue, quarter to date", f"=Regional!B{qtd_t}", MONEY),
        ("Quota to date, quarter", f"=Regional!C{qtd_t}", MONEY),
        ("Attainment, quarter to date", f"=Regional!D{qtd_t}", PCT),
        ("Attainment, year to date", f"=Regional!D{ytd_t}", PCT),
        ("Avg order value, quarter to date", f"=Regional!G{qtd_t}", MONEY),
        ("Growth vs last year, quarter to date", f"=Regional!I{qtd_t}", PCT),
    ]
    for r, (label, formula, fmt) in enumerate(kpis, start=14):
        ws.cell(row=r, column=1, value=label)
        _put(ws, r, 2, formula, fmt, bold=True)
    _attainment_colors(ws, "B17:B18")

    top = 14 + len(kpis) + 2
    _section(ws, top, "Quarter to date by region", 6)
    _header(ws, top + 1, ["Region", "Revenue", "Quota to date", "Attainment", "Gap to quota", "Growth vs PY"])
    sources = [layout["qtd"]["rows"][reg] for reg in regions] + [qtd_t]
    for k, src in enumerate(sources):
        r = top + 2 + k
        for c, (col, fmt) in enumerate(zip("ABCDEI", (None, MONEY, MONEY, PCT, MONEY, PCT)), start=1):
            _put(ws, r, c, f"=Regional!{col}{src}", fmt)
    last = top + 1 + len(sources)
    _total_row(ws, last, 1, 6)
    _attainment_colors(ws, f"D{top + 2}:D{last}")

    ctop = last + 2
    _section(ws, ctop, f"Key callouts for the week ending {info.as_of:%b %d, %Y}", 6)
    for k, text in enumerate(callouts, start=1):
        r = ctop + k
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        cell = ws.cell(row=r, column=1, value=f"• {text}")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 13.5 * (1 + (len(text) + 2) // 125)  # Excel won't auto-fit merged rows
    ws.cell(row=ctop + len(callouts) + 2, column=1,
            value="Callouts are written by the Python refresh for the as-of date above; rerun it to update them."
            ).font = _font(italic=True, color=GREY, size=9)
    _widths(ws, {"A": 36, "B": 16, "C": 16, "D": 14, "E": 16, "F": 14})


def _write_definitions(ws, info: RunInfo) -> None:
    _title(ws, "Definitions and run details")
    _header(ws, 3, ["Term", "Definition"])
    terms = [
        ("Revenue", "Sum of Sales for order lines dated inside the window."),
        ("Orders", "Distinct orders. The first line of each order carries Order Count = 1, so sums count each order once."),
        ("Average order value", "Revenue ÷ Orders."),
        ("Quota", "Monthly target by region. Default: same month last year × (1 + growth on the Summary sheet). "
                  "Rows marked 'override' come from the quota file and ignore the growth input."),
        ("Quota to date", "Each month's quota × (days of that month inside the window ÷ days in the month)."),
        ("Attainment", "Revenue ÷ Quota to date: pace against target so far."),
        ("Gap to quota", "Revenue − Quota to date. Negative means behind pace."),
        ("Growth vs PY", "Revenue ÷ revenue for the same dates last year − 1. For the week, last year's "
                         "week is 52 weeks earlier so the weekdays line up."),
        ("Week", "Monday to Sunday, labelled by the Sunday. Last year's week is 52 weeks earlier."),
        ("Mix shift", "Change in a category's share of revenue vs last year, in percentage points."),
    ]
    for r, (term, text) in enumerate(terms, start=4):
        ws.cell(row=r, column=1, value=term).font = _font(bold=True)
        ws.cell(row=r, column=2, value=text).alignment = Alignment(wrap_text=True, vertical="top")

    top = 4 + len(terms) + 1
    _header(ws, top, ["Run detail", "Value"])
    details = [
        ("Source file", info.source),
        ("Rows loaded", f"{info.rows_loaded:,}"),
        ("Rows dropped (missing date, sales, order ID, region or category)", f"{info.rows_dropped:,}"),
        ("Data covers", f"{info.first_date:%b %d, %Y} to {info.last_date:%b %d, %Y}"),
        ("As-of date of the refresh run", f"{info.as_of:%b %d, %Y}"),
        ("Quota method", info.quota_method),
        ("Built", f"{datetime.now():%Y-%m-%d %H:%M} by sales_pack v{__version__}"),
    ]
    for r, (label, value) in enumerate(details, start=top + 1):
        ws.cell(row=r, column=1, value=label).font = _font(bold=True)
        ws.cell(row=r, column=2, value=value).alignment = Alignment(wrap_text=True, vertical="top")
    _widths(ws, {"A": 34, "B": 100})
